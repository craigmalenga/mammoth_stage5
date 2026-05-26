"""Extract CCJ / credit-search data from Mammoth Stage 5 PDFs into an Excel file.

Usage:
    python3 extract_stage5.py <pdf_dir> <output_xlsx> [--limit N]
"""

import argparse
import re
import sys
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font

# Patterns are anchored line-by-line via re.MULTILINE.
# Each label phrase is unique within the form, so simple ".*?(value)\s*$" is safe.
PATTERNS = {
    "account_no":  re.compile(r"^ACCOUNT NO:\s*(\d+)\s*$",                                       re.MULTILINE),
    "unsatisfied": re.compile(r"^#\s*unsatisfied CCJs.*?(\d+)\s*$",                              re.MULTILINE | re.IGNORECASE),
    "satisfied":   re.compile(r"^#\s*satisfied CCJs.*?(\d+)\s*$",                                re.MULTILINE | re.IGNORECASE),
    "searches_3m": re.compile(r"^#\s*credit searches last 3 months.*?(\d+)\s*$",                 re.MULTILINE | re.IGNORECASE),
    "searches_12m":re.compile(r"^#\s*credit searches last 12 months.*?(\d+)\s*$",                re.MULTILINE | re.IGNORECASE),
    "score":       re.compile(r"^Resulting score.*?(\d+(?:\.\d+)?)\s*$",                          re.MULTILINE | re.IGNORECASE),
}


def extract_one(pdf_path: Path) -> dict:
    """Extract the 6 fields from one PDF. Returns dict with fields + 'error' if any."""
    result = {"file": pdf_path.name, "error": None}
    try:
        with pdfplumber.open(pdf_path) as pdf:
            text = "\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception as e:
        result["error"] = f"open/extract failed: {e}"
        return result

    missing = []
    for key, pat in PATTERNS.items():
        m = pat.search(text)
        if m:
            result[key] = m.group(1)
        else:
            result[key] = None
            missing.append(key)
    if missing:
        result["error"] = "missing: " + ",".join(missing)
    return result


def to_number(s, kind):
    if s is None:
        return None
    if kind == "int":
        return int(s)
    if kind == "float":
        return float(s)
    return s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pdf_dir")
    ap.add_argument("output_xlsx")
    ap.add_argument("--limit", type=int, default=None)
    args = ap.parse_args()

    pdf_dir = Path(args.pdf_dir)
    pdfs = sorted(pdf_dir.glob("*.pdf"))
    if args.limit:
        pdfs = pdfs[: args.limit]

    print(f"Processing {len(pdfs)} PDFs from {pdf_dir} ...", flush=True)

    rows = []
    errors = []
    for i, p in enumerate(pdfs, 1):
        r = extract_one(p)
        if r.get("error"):
            errors.append((p.name, r["error"]))
        rows.append(r)
        if i % 50 == 0 or i == len(pdfs):
            print(f"  ... {i}/{len(pdfs)} done", flush=True)

    # Build workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Stage5 CCJs"

    headers = [
        "Lead ID",
        "# unsatisfied CCJs",
        "# satisfied CCJs",
        "# credit searches 3m",
        "# credit searches 12m",
        "Score",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in rows:
        ws.append([
            to_number(r.get("account_no"), "int"),
            to_number(r.get("unsatisfied"), "int"),
            to_number(r.get("satisfied"), "int"),
            to_number(r.get("searches_3m"), "int"),
            to_number(r.get("searches_12m"), "int"),
            to_number(r.get("score"), "float"),
        ])

    # Number formats
    int_cols = ["A", "B", "C", "D", "E"]
    last_row = ws.max_row
    for col in int_cols:
        for row in range(2, last_row + 1):
            ws[f"{col}{row}"].number_format = "0"
    for row in range(2, last_row + 1):
        ws[f"F{row}"].number_format = "0.0"

    # Column widths
    widths = {"A": 12, "B": 22, "C": 22, "D": 22, "E": 22, "F": 8}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # Errors sheet (only if any)
    if errors:
        es = wb.create_sheet("Errors")
        es.append(["File", "Error"])
        for f, e in errors:
            es.append([f, e])

    wb.save(args.output_xlsx)

    print(f"\nSaved {args.output_xlsx}")
    print(f"  rows written: {len(rows)}")
    print(f"  errors      : {len(errors)}")
    if errors:
        print("  first 5 errors:")
        for f, e in errors[:5]:
            print(f"    {f}: {e}")


if __name__ == "__main__":
    main()
